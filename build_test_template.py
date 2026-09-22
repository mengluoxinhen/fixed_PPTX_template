"""Generate the test PPTX template and the test product image.

Run once before rendering:

    python build_test_template.py

Creates:
    templates/sales_report_template.pptx  (4 slides, typed {{text:/{{image:}}
                                            placeholders incl. repeats)
    assets/product.png                    (Pillow-generated mockup)
    assets/logo.jpg                       (Pillow-generated .jpg asset)
"""
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE_PATH = ROOT / "templates" / "sales_report_template.pptx"
IMAGE_PATH = ROOT / "assets" / "product.png"
LOGO_PATH = ROOT / "assets" / "logo.jpg"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xF2, 0x99, 0x38)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
DARK = RGBColor(0x2B, 0x2F, 0x38)
GRAY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
CARD_BLUE = RGBColor(0x2E, 0x5E, 0xAA)


def add_text(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT):
    """lines: list of (text, size_pt, bold, color, [font_name])."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color, *rest) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = rest[0] if rest else "Calibri"
    return box


def add_rect(slide, shape_type, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(
        shape_type, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def add_image_region(slide, left, top, width, height, key, font_size=12):
    """A gray box whose whole text is one {{image:key}} = the image region."""
    ph = add_rect(slide, MSO_SHAPE.RECTANGLE, left, top, width, height,
                  RGBColor(0xE3, 0xE8, 0xEF), line=RGBColor(0xB8, 0xC2, 0xD1))
    tf = ph.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"{{{{image:{key}}}}}"
    run.font.size = Pt(font_size)
    run.font.color.rgb = GRAY
    run.font.name = "Calibri"
    return ph


def build_cover(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY

    add_rect(slide, MSO_SHAPE.RECTANGLE, 0.9, 2.35, 1.6, 0.06, ORANGE)

    add_text(slide, 0.9, 2.6, 11.5, 1.2, [
        ("{{text:title}}", 44, True, WHITE),
    ])
    add_text(slide, 0.9, 3.75, 11.5, 0.7, [
        ("{{text:subtitle}}", 22, False, RGBColor(0xC9, 0xD6, 0xE8)),
    ])
    add_text(slide, 0.9, 4.55, 11.5, 0.5, [
        ("{{text:report_date}}", 15, False, ORANGE),
    ])
    add_text(slide, 0.9, 6.6, 11.5, 0.5, [
        ("{{text:company_name}}", 13, False, RGBColor(0x9A, 0xAC, 0xC6)),
    ])
    # Second occurrence of {{image:product_image}} (square region, slide 1)
    add_image_region(slide, 10.9, 4.95, 1.6, 1.6, "product_image", font_size=9)


def build_workspace_assets(slide):
    """Slide 4: stem-keyed image regions for workspace mode.

    product/logo assets resolve from a UUID workspace (or example.json in
    direct mode); missing_image has no asset anywhere and must stay
    unchanged after rendering.
    """
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    add_text(slide, 0.9, 0.55, 11.5, 0.8, [
        ("Workspace Asset Test", 28, True, NAVY),
    ])
    add_rect(slide, MSO_SHAPE.RECTANGLE, 0.95, 1.25, 1.2, 0.05, ORANGE)

    add_image_region(slide, 0.9, 1.9, 2.6, 2.6, "product", font_size=11)
    add_image_region(slide, 3.9, 1.9, 3.6, 1.2, "logo", font_size=11)
    # Repeat of the same image key with a different (tall) region shape
    add_image_region(slide, 7.9, 1.9, 1.2, 3.4, "product", font_size=9)
    # No asset exists for this key: placeholder must survive rendering
    add_image_region(slide, 9.6, 1.9, 2.2, 2.2, "missing_image", font_size=11)


def build_metrics(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT

    add_text(slide, 0.9, 0.55, 11.5, 0.8, [
        ("Key Business Metrics", 28, True, NAVY),
    ])
    add_rect(slide, MSO_SHAPE.RECTANGLE, 0.95, 1.25, 1.2, 0.05, ORANGE)

    # Repeated text occurrence across slides: {{text:title}} again (slide 2)
    add_text(slide, 0.9, 1.45, 11.5, 0.45, [
        ("Report: {{text:title}}", 13, False, GRAY),
    ])

    cards = [
        ("Sales", "{{text:sales}}", CARD_BLUE),
        ("Growth", "{{text:growth}}", GREEN),
        ("Orders", "{{text:orders}}", CARD_BLUE),
        ("Customers", "{{text:customers}}", CARD_BLUE),
    ]
    card_w, card_h, gap = 2.75, 2.6, 0.32
    left0 = (13.333 - (card_w * 4 + gap * 3)) / 2
    for i, (label, var, value_color) in enumerate(cards):
        left = left0 + i * (card_w + gap)
        add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, 2.0, card_w, card_h, WHITE)
        add_rect(slide, MSO_SHAPE.RECTANGLE, left + 0.25, 2.45, 0.5, 0.05, ORANGE)
        add_text(slide, left + 0.25, 2.1, card_w - 0.5, 0.4, [
            (label.upper(), 12, True, GRAY),
        ])
        add_text(slide, left + 0.25, 2.75, card_w - 0.5, 1.0, [
            (var, 24, True, value_color),
        ])

    add_text(slide, 0.9, 5.6, 11.5, 0.5, [
        ("Total reported sales: {{text:sales}} year to date.", 14, False, GRAY),
    ])


def build_product(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    add_text(slide, 0.9, 0.55, 11.5, 0.8, [
        ("Product Showcase", 28, True, NAVY),
    ])
    add_rect(slide, MSO_SHAPE.RECTANGLE, 0.95, 1.25, 1.2, 0.05, ORANGE)

    # Left text column
    add_text(slide, 0.9, 2.0, 5.6, 0.8, [
        ("{{text:product_name}}", 26, True, DARK),
    ])
    add_text(slide, 0.9, 2.85, 5.6, 1.4, [
        ("{{text:product_description}}", 15, False, GRAY),
    ])
    add_text(slide, 0.9, 4.4, 5.6, 0.7, [
        ("Price", 12, True, GRAY),
        ("{{text:product_price}}", 24, True, ORANGE),
    ])
    add_text(slide, 0.9, 5.5, 5.6, 0.7, [
        ("Month-over-month growth", 12, True, GRAY),
        ("{{text:product_growth}}", 20, True, GREEN),
    ])

    # Right image region (placeholder replaced by the renderer)
    add_image_region(slide, 7.1, 1.9, 5.3, 4.6, "product_image", font_size=14)

    # Repeated text occurrence across slides: {{text:title}} again (slide 3)
    add_text(slide, 0.9, 6.95, 11.5, 0.4, [
        ("{{text:title}}", 11, False, GRAY),
    ])


def build_template():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for builder in (build_cover, build_metrics, build_product, build_workspace_assets):
        builder(prs.slides.add_slide(blank))

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(TEMPLATE_PATH)
    print(f"Template written: {TEMPLATE_PATH}")


def load_font(size):
    for name in ("arialbd.ttf", "arial.ttf", "segoeui.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def build_product_image():
    """Simple device mockup on a soft gradient background."""
    W, H = 1200, 1040
    img = Image.new("RGB", (W, H))
    draw = ImageDraw.Draw(img)

    top_c, bot_c = (0xDD, 0xE7, 0xF5), (0xF7, 0xFA, 0xFD)
    for y in range(H):
        t = y / H
        draw.line(
            [(0, y), (W, y)],
            fill=tuple(int(a + (b - a) * t) for a, b in zip(top_c, bot_c)),
        )

    # Terminal body
    body = (330, 190, 870, 700)
    draw.rounded_rectangle(body, radius=36, fill=(0x2B, 0x33, 0x42))
    # Screen
    draw.rounded_rectangle((365, 225, 835, 620), radius=18, fill=(0x14, 0x1A, 0x24))
    draw.rounded_rectangle((400, 265, 720, 305), radius=12, fill=(0xF2, 0x99, 0x38))
    for i, bar_w in enumerate((320, 250, 290, 210)):
        draw.rounded_rectangle((400, 345 + i * 52, 400 + bar_w, 375 + i * 52),
                               radius=8, fill=(0x3D, 0x8F, 0xC9))
    # Stand
    draw.rectangle((560, 700, 640, 780), fill=(0x39, 0x42, 0x52))
    draw.rounded_rectangle((450, 775, 750, 815), radius=16, fill=(0x2B, 0x33, 0x42))

    draw.text((W // 2, 900), "Smart Office Terminal X1", font=load_font(44),
              fill=(0x2B, 0x33, 0x42), anchor="mm")
    draw.text((W // 2, 960), "product placeholder image", font=load_font(26),
              fill=(0x8A, 0x94, 0xA6), anchor="mm")

    IMAGE_PATH.parent.mkdir(parents=True, exist_ok=True)
    img.save(IMAGE_PATH)
    print(f"Image written: {IMAGE_PATH}")


def build_logo_image():
    """Small .jpg asset to prove stem resolution is extension-independent."""
    img = Image.new("RGB", (800, 600), (0x1F, 0x3A, 0x5F))
    draw = ImageDraw.Draw(img)
    draw.ellipse((250, 130, 550, 430), outline=(0xF2, 0x99, 0x38), width=28)
    draw.rectangle((360, 240, 440, 320), fill=(0xF2, 0x99, 0x38))
    draw.text((400, 500), "EXAMPLE LOGO", font=load_font(48),
              fill=(0xFF, 0xFF, 0xFF), anchor="mm")
    LOGO_PATH.parent.mkdir(parents=True, exist_ok=True)
    img.save(LOGO_PATH)
    print(f"Image written: {LOGO_PATH}")


if __name__ == "__main__":
    build_template()
    build_product_image()
    build_logo_image()
