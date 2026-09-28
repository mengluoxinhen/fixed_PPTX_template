"""Debug tool: dump everything the renderer reads from a PPTX template.

Run directly, with no command-line arguments:

    python inspect_template.py

All knobs are the constants below — edit them in code for each debugging
session:

    TEMPLATE   path to the .pptx template (required)
    DATA       direct-mode JSON data file, or None
    WORKSPACE  UUID4 workspace directory, or None (wins over DATA)
    OUT        where the JSON dump is written

The dump contains:
  - every text container (shape or table cell): geometry, full text,
    recognized {{text:key}} / {{image:key}} tokens, malformed tokens;
  - the parser's own view (parse_slides): text/image placeholders and
    per-occurrence image regions;
  - the data source actually available to the renderer;
  - a binding check: which placeholders would be rendered and which
    would REMAIN unchanged because their key is missing from the data.

Read-only: this script never renders and never modifies the template.
"""
import json
import sys
from pathlib import Path

from pptx import Presentation

from src.template_parser import (extract_placeholders, find_malformed,
                                 iter_text_shapes, parse_slides)

# ---- Debug knobs: edit these, then run `python inspect_template.py` ----
TEMPLATE = "templates/sales_report_template.pptx"
DATA = "data/example.json"   # None to skip direct-mode data
WORKSPACE = None             # e.g. "workspace/550e8400-e29b-41d4-a716-446655440000"; beats DATA
OUT = "debug/template_dump.json"
# ------------------------------------------------------------------------

EMU_PER_INCH = 914400
KEY_CHARSET_NOTE = "keys must match [A-Za-z0-9_]+"


def emu_box(shape):
    left, top = getattr(shape, "left", None), getattr(shape, "top", None)
    if left is None or top is None:
        return None
    return [int(left), int(top), int(shape.width), int(shape.height)]


def inches(box):
    return None if box is None else [round(v / EMU_PER_INCH, 4) for v in box]


def container_label(holder):
    name = getattr(holder, "name", None)
    if name:
        return name
    # Table cells carry no shape name/geometry.
    return "table cell" if getattr(holder, "left", None) is None else str(holder)


def iter_table_shapes(shapes):
    """Yield graphic-frame shapes that host a table, recursing into groups."""
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from iter_table_shapes(shape.shapes)
        elif getattr(shape, "has_table", False):
            yield shape


def describe_text_containers(prs):
    """Walk shapes AND table cells; report text + recognized tokens."""
    slides_out = []
    for slide_index, slide in enumerate(prs.slides):
        shapes = []
        for shape in iter_text_shapes(slide.shapes):
            text = shape.text_frame.text
            box = emu_box(shape)
            shapes.append({
                "container": container_label(shape),
                "kind": "shape",
                "geometry_emu": box,
                "geometry_inches": inches(box),
                "text": text,
                "text_keys": [k for t, k in extract_placeholders(text) if t == "text"],
                "image_keys": [k for t, k in extract_placeholders(text) if t == "image"],
                "malformed": find_malformed(text),
            })
        cells = []
        for host in iter_table_shapes(slide.shapes):
            for r, row in enumerate(host.table.rows):
                for c, cell in enumerate(row.cells):
                    text = cell.text_frame.text
                    cells.append({
                        "container": f"{container_label(host)}[row {r}, col {c}]",
                        "kind": "table_cell",
                        "text": text,
                        "text_keys": [k for t, k in extract_placeholders(text)
                                      if t == "text"],
                        "image_keys": [k for t, k in extract_placeholders(text)
                                       if t == "image"],
                        "malformed": find_malformed(text),
                        "note": "table cells carry TEXT placeholders only; "
                                "{{image:key}} inside a cell stays unchanged",
                    })
        slides_out.append({"slide": slide_index, "shapes": shapes, "table_cells": cells})
    return slides_out


def describe_parsed(prs):
    text_phs, image_phs, malformed = parse_slides(prs, {})
    return {
        "text_placeholders": [
            {
                "slide": ph.slide_index,
                "container": container_label(ph.shape),
                "keys": ph.keys,
            }
            for ph in text_phs
        ],
        "image_placeholders": [
            {
                "slide": ph.slide_index,
                "container": container_label(ph.shape),
                "key": ph.key,
                "region_emu": [int(v) for v in ph.region],
                "region_inches": inches([int(v) for v in ph.region]),
            }
            for ph in image_phs
        ],
        "malformed_tokens": sorted(set(malformed)),
    }


def binding_check(prs, text_data, image_lookup_keys, mode):
    """Which placeholders would render vs remain unchanged for this data."""
    required_text, required_image = set(), set()
    for slide_info in describe_text_containers(prs):
        for entry in slide_info["shapes"] + slide_info["table_cells"]:
            required_text.update(entry["text_keys"])
            required_image.update(entry["image_keys"])
    _, image_phs, _ = parse_slides(prs, {})
    # Only whole-shape {{image:key}} regions can actually render.
    required_image = {ph.key for ph in image_phs}

    text_present = sorted(required_text & set(text_data))
    image_present = sorted(required_image & set(image_lookup_keys))
    return {
        "mode": mode,
        "text_keys": {"required": sorted(required_text),
                      "will_render": text_present,
                      "missing_will_remain": sorted(required_text - set(text_data))},
        "image_keys": {"required": sorted(required_image),
                       "will_render": image_present,
                       "missing_will_remain": sorted(required_image - set(image_lookup_keys))},
    }


def preview_value(value):
    s = str(value)
    return s if len(s) <= 120 else s[:117] + "..."


def main():
    template_path = Path(TEMPLATE)
    if not template_path.is_file():
        sys.exit(f"Template not found: {template_path} (edit TEMPLATE in this file)")

    if WORKSPACE:
        from src.workspace_loader import load_workspace

        workspace = load_workspace(WORKSPACE)
        text_data = workspace["data"]
        image_available = workspace["image_paths"]
        data_source = {
            "mode": "workspace",
            "workspace": workspace["path"],
            "content_json_keys": sorted(text_data),
            "content_json_values_preview": {k: preview_value(v) for k, v in text_data.items()},
            "image_paths": workspace["image_paths"],
        }
    elif DATA:
        with open(DATA, "r", encoding="utf-8") as f:
            text_data = json.load(f)
        image_available = text_data  # direct mode: image keys live in the same dict
        data_source = {
            "mode": "direct",
            "data_file": str(DATA),
            "data_keys": sorted(text_data),
            "data_values_preview": {k: preview_value(v) for k, v in text_data.items()},
        }
    else:
        text_data = {}
        image_available = {}
        data_source = {"mode": "none",
                       "note": "DATA and WORKSPACE are both None; "
                               "binding check marks everything missing"}

    prs = Presentation(str(template_path))
    dump = {
        "tool": "inspect_template.py",
        "template": str(template_path),
        "slide_count": len(prs.slides._sldIdLst),
        "slide_size_emu": [int(prs.slide_width), int(prs.slide_height)],
        "key_charset_note": KEY_CHARSET_NOTE,
        "slides": describe_text_containers(prs),
        "parsed_placeholders": describe_parsed(prs),
        "data_source": data_source,
        "binding_check": binding_check(prs, text_data, image_available,
                                       data_source["mode"]),
    }

    out_path = Path(OUT)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(dump, f, ensure_ascii=False, indent=2)
    print(f"Template dump written: {out_path}")

    bc = dump["binding_check"]
    print(f"  text : {len(bc['text_keys']['will_render'])} render, "
          f"{len(bc['text_keys']['missing_will_remain'])} missing "
          f"{bc['text_keys']['missing_will_remain'] or ''}")
    print(f"  image: {len(bc['image_keys']['will_render'])} render, "
          f"{len(bc['image_keys']['missing_will_remain'])} missing "
          f"{bc['image_keys']['missing_will_remain'] or ''}")
    if bc["image_keys"]["missing_will_remain"] and data_source["mode"] == "workspace":
        print("  hint : workspace image keys are FILENAME STEMS without extension "
              "(product.png -> key 'product'); the key is NOT 'product.png'.")


if __name__ == "__main__":
    main()
