"""Replace {{image:key}} placeholder shapes with real pictures.

Only the template declares a region as an image: the placeholder type is
explicit, so values are never inspected for extensions or path shape.

The placeholder shape defines the region (position + size). The renderer
deletes the placeholder and inserts the image using "contain" fitting:
the image is uniformly scaled (aspect ratio preserved, never stretched)
to the largest size that fits ENTIRELY inside the region, centered on the
region. No content is cropped; if the aspect ratios differ, the image
simply does not touch the region's shorter edges.

The picture is re-inserted at the placeholder's original z-order position
so bottom-layer placeholders stay bottom-layer after replacement.
"""
from PIL import Image

from pptx.util import Emu


def _contain_size(img_w, img_h, region_w, region_h):
    """Largest (w, h) with the image's aspect ratio that fits the region."""
    scale = min(region_w / img_w, region_h / img_h)
    return int(img_w * scale), int(img_h * scale)


def render_images(prs, data, resolve_path, placeholders=None):
    """Replace each {{image:key}} placeholder with its picture.

    resolve_path(value) must return an existing image file path.
    Returns list of (key, used_path). Keys missing from data are skipped
    so the placeholder remains visible.
    """
    if placeholders is None:
        from .template_parser import parse_slides

        _, placeholders, _ = parse_slides(prs, data)

    used = []
    for ph in placeholders:
        if ph.key not in data:
            continue
        path = resolve_path(data[ph.key])
        left, top, width, height = ph.region
        slide = prs.slides[ph.slide_index]

        with Image.open(path) as im:
            img_w, img_h = im.size
        pic_w, pic_h = _contain_size(img_w, img_h, width, height)
        pic_left = left + (width - pic_w) // 2
        pic_top = top + (height - pic_h) // 2

        # Remove the placeholder at its z-order slot, then insert the
        # picture back into that same slot (bottom-layer stays bottom).
        element = ph.shape._element
        parent = element.getparent()
        sp_tree = slide.shapes._spTree
        z_index = sp_tree.index(element) if parent is sp_tree else None
        parent.remove(element)

        pic = slide.shapes.add_picture(
            path, Emu(pic_left), Emu(pic_top), width=Emu(pic_w), height=Emu(pic_h)
        )
        if z_index is not None:
            pic_element = pic._element
            sp_tree.remove(pic_element)
            sp_tree.insert(z_index, pic_element)

        used.append((ph.key, path))

    return used
