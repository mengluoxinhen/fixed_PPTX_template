"""Generic template renderer: PPTX template + JSON data -> populated PPTX.

Business-agnostic: the only contract is an explicit {{text:key}} /
{{image:key}} placeholder in the template matched against keys in the
data dict. The template owns the type; the data owns the value.
"""
import json
from pathlib import Path

from pptx import Presentation

from .template_parser import MalformedPlaceholderError, parse_slides
from .text_renderer import render_text
from .image_renderer import render_images


class MissingAssetError(FileNotFoundError):
    pass


def load_data(data_path):
    """Load JSON data and record the base dir used to resolve asset paths."""
    data_path = Path(data_path)
    with open(data_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return data, data_path.parent


def render(template_path, data, output_path, base_dirs=(".",)):
    """Populate a PPTX template with data and write the result.

    Args:
        template_path: source .pptx template (never modified).
        data: dict of key -> value.
        output_path: destination .pptx.
        base_dirs: directories searched, in order, to resolve relative
            image paths found in the data.

    Raises:
        MalformedPlaceholderError: if the template contains any {{...}}
            token that is not exactly {{text:key}} or {{image:key}}.
    """
    prs = Presentation(template_path)

    _, image_placeholders, malformed = parse_slides(prs, data)
    if malformed:
        raise MalformedPlaceholderError(malformed)

    def resolve_path(value):
        p = Path(str(value))
        if p.is_absolute() and p.exists():
            return str(p)
        for base in base_dirs:
            candidate = Path(base) / value
            if candidate.exists():
                return str(candidate)
        raise MissingAssetError(f"Image asset not found: {value}")

    # Images first: placeholders removed, region geometry reused.
    used_images = render_images(prs, data, resolve_path, image_placeholders)
    # Then remaining text placeholders.
    replaced = render_text(prs, data)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)

    return {
        "text_keys": sorted(set(replaced)),
        "image_keys": sorted(key for key, _ in used_images),
    }
