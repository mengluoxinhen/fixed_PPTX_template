"""UUID workspace adapter: one PPT generation task per UUID4 folder.

    <workspace>/<uuid4>/
        content.json     text resources ONLY
        product.png      image resources, resolved by filename stem
        logo.jpg         via {{image:product}} / {{image:logo}} in the template

Supported image formats: .png / .jpg / .jpeg. Other files (including .webp)
are not treated as image resources.

Strict namespace separation — the binding identity is (type, key):

    {{text:key}}   ->  content.json[key]            (text namespace)
    {{image:key}}  ->  workspace/<key>.<ext>        (image namespace)

A key may exist in both namespaces at once (e.g. content.json["logo"] and
logo.jpg); neither lookup interferes with the other. Image files are never
consulted for text and content.json values are never consulted for images.

This module is a loading/composition adapter only: it adds no workspace
logic to the core renderer files (renderer.py, template_parser.py,
text_renderer.py, image_renderer.py). The supported image extension list
lives here (and in the template's explicit {{image:}} type), never in the
renderer.
"""
import json
import shutil
import uuid
from pathlib import Path

from pptx import Presentation

from .image_renderer import render_images
from .template_parser import MalformedPlaceholderError, parse_slides
from .text_renderer import render_text

# Centralized supported image formats; order = priority on stem conflict.
# Unsupported files (e.g. .webp) are ignored by discovery, so the matching
# {{image:key}} placeholder simply remains unchanged.
SUPPORTED_IMAGE_EXTS = (".png", ".jpg", ".jpeg")

CONTENT_FILENAME = "content.json"


class WorkspaceError(Exception):
    """Raised for invalid or unreadable workspaces."""


def validate_workspace_dir(workspace_path):
    """Ensure the folder itself is a directory named by a valid UUID4."""
    path = Path(workspace_path)
    if not path.is_dir():
        raise WorkspaceError(f"Workspace directory not found: {path}")
    try:
        parsed = uuid.UUID(path.name)
    except ValueError:
        raise WorkspaceError(
            f"Workspace folder name is not a valid UUID: {path.name!r}"
        ) from None
    if parsed.version != 4:
        raise WorkspaceError(f"Workspace folder name is not a UUID4: {path.name!r}")
    return path


def discover_images(workspace_dir):
    """Map image filename stems to files located directly in the workspace.

    Only regular files with a supported extension are considered; other
    files and subdirectories are ignored. When the same stem exists with
    several extensions, the first one in SUPPORTED_IMAGE_EXTS order wins.
    Every returned path is guaranteed to stay inside the workspace.
    """
    root = Path(workspace_dir).resolve()
    images = {}
    for ext in SUPPORTED_IMAGE_EXTS:  # priority order: earlier ext wins
        for entry in sorted(root.glob(f"*{ext}")):
            if entry.is_file() and entry.stem not in images:
                images[entry.stem] = str(entry.resolve())
    # Defense in depth: never hand back a path outside the workspace.
    return {
        stem: p for stem, p in images.items()
        if Path(p).is_relative_to(root)
    }


def load_workspace(workspace_path):
    """Validate a UUID workspace and load its two independent namespaces.

    Returns a plain dict, renderer-ready:

        {
            "path":        str,  absolute workspace directory,
            "data":        dict, content.json text values (no image paths),
            "image_paths": dict, image stem -> absolute file path,
        }

    Missing resources are simply absent from the returned dicts, which the
    renderers handle by leaving the matching placeholders unchanged.
    """
    path = validate_workspace_dir(workspace_path).resolve()

    content_file = path / CONTENT_FILENAME
    if not content_file.is_file():
        raise WorkspaceError(f"Workspace is missing content.json: {path}")
    with open(content_file, "r", encoding="utf-8") as f:
        content = json.load(f)
    if not isinstance(content, dict):
        raise WorkspaceError("content.json must contain a JSON object")

    return {
        "path": str(path),
        "data": content,
        "image_paths": discover_images(path),
    }


def render_workspace(template_path, workspace, output_path):
    """Render using the existing primitives with strict (type, key) bindings.

    A flat data dict cannot carry two different values for one key, so
    instead of merging namespaces this adapter composes the existing,
    unmodified render pipeline per binding type:

        parse_slides   -> classify placeholders by declared type
        render_images  -> reads ONLY workspace["image_paths"]
        render_text    -> reads ONLY workspace["data"]
    """
    prs = Presentation(template_path)

    _, image_placeholders, malformed = parse_slides(prs, workspace["data"])
    if malformed:
        raise MalformedPlaceholderError(malformed)

    # Image stems are already validated absolute paths inside the workspace.
    def resolve_path(value):
        return str(value)

    used_images = render_images(
        prs, workspace["image_paths"], resolve_path, image_placeholders
    )
    replaced = render_text(prs, workspace["data"])

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)

    return {
        "text_keys": sorted(set(replaced)),
        "image_keys": sorted(key for key, _ in used_images),
    }


def cleanup_workspace(workspace_path):
    """Explicitly delete a UUID workspace (opt-in; never called implicitly).

    Refuses anything whose folder name is not a valid UUID4, so a typo
    cannot wipe an arbitrary directory.
    """
    path = validate_workspace_dir(workspace_path)
    shutil.rmtree(path)
    return path
